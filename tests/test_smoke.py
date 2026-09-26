from __future__ import annotations

import asyncio
import importlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile


class StudioSmokeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp = tempfile.TemporaryDirectory()
        os.environ["H3_STUDIO_DATA_DIR"] = cls.temp.name
        os.environ["H3_COMFYUI_INPUT_DIR"] = str(Path(cls.temp.name) / "comfyui" / "input")
        os.environ["H3_COMFYUI_OUTPUT_DIR"] = str(Path(cls.temp.name) / "comfyui" / "output")
        import app.main
        cls.main = importlib.reload(app.main)
        cls.main.init_db()

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp.cleanup()

    def test_01_queue_prompt_completion_and_delete(self) -> None:
        folder = Path(self.temp.name) / "fixtures"
        folder.mkdir(exist_ok=True)
        source = folder / "source.mp4"
        person = folder / "person.jpg"
        product = folder / "product.png"
        source.write_bytes(b"video")
        person.write_bytes(b"image")
        product.write_bytes(b"image")

        task_id = self.main.insert_task(
            prompt="Keep the original background.", source_video=source,
            reference_images=[person, product], duration=5.25,
            aspect_ratio="16:9", seed=42, display_name="smoke.mp4",
        )
        row = self.main.claim_next_task()
        self.assertEqual(row["id"], task_id)
        prompt = self.main.build_h3_prompt(row)
        self.assertIn("<Video 1>", prompt)
        self.assertIn("<Picture 1>", prompt)
        self.assertIn("<Picture 2>", prompt)
        self.assertIn("change only the visible facial identity", prompt)
        self.assertIn("main model must never disappear", prompt)
        self.assertIn("replace only the pixels belonging to the original product", prompt)

        def fake_http(method, url, payload=None, timeout=60):
            if url.endswith("/prompt"):
                workflow = payload["prompt"]
                sampler = workflow["136"]["inputs"]
                self.assertIn("ref_videos.ref_video_0", sampler)
                self.assertIn("ref_images.ref_image_0", sampler)
                self.assertIn("ref_images.ref_image_1", sampler)
                self.assertEqual(sampler["width"], 672)
                self.assertEqual(sampler["height"], 384)
                return {"prompt_id": "comfy-1"}
            if "/history/comfy-1" in url:
                return {"comfy-1": {
                    "status": {"status_str": "success"},
                    "outputs": {"92": {"images": [{
                        "filename": "result.mp4", "subfolder": "h3_studio", "type": "output"
                    }]}},
                }}
            return {}

        def fake_download(_, destination):
            destination.write_bytes(b"fake-mp4")

        with patch.object(self.main, "http_json", side_effect=fake_http), patch.object(
            self.main, "download_comfy_output", side_effect=fake_download
        ), patch.object(self.main, "open_comfy_websocket", return_value=None):
            self.main.process_task(row)

        with self.main.connect_db() as conn:
            completed = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["stage"], "生成完成")
        self.assertEqual((completed["output_width"], completed["output_height"]), (672, 384))
        self.assertTrue(Path(completed["output_path"]).is_file())
        result = self.main.delete_task(task_id)
        self.assertTrue(result["deleted"])
        self.assertFalse(Path(completed["output_path"]).exists())

    def test_02_excel_import_and_template(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "任务"
        sheet.append([
            "upload_video_url", "reference_image_urls", "reference_roles", "product_box",
            "prompt", "aspect_ratio", "quality", "seed",
        ])
        sheet.append([
            "https://media.example.com/bulk.mp4",
            "https://media.example.com/face.jpg",
            "face", "", "自然光", "9:16", "low", 99,
        ])
        excel_data = io.BytesIO()
        workbook.save(excel_data)
        excel_data.seek(0)

        excel = UploadFile(filename="tasks.xlsx", file=excel_data)
        with patch.object(self.main, "probe_video_duration", return_value=4.5):
            result = asyncio.run(self.main.import_excel_tasks(excel=excel))
        self.assertEqual(result["created"], 1)
        self.assertEqual(result["errors"], [])

        response = self.main.excel_template()

        async def consume() -> bytes:
            parts = []
            async for part in response.body_iterator:
                parts.append(part)
            return b"".join(parts)

        template_bytes = asyncio.run(consume())
        template = load_workbook(io.BytesIO(template_bytes))
        self.assertEqual(template.sheetnames, ["任务", "字段说明"])
        self.assertEqual(template["任务"]["A1"].value, "upload_video_url")
        self.assertEqual(template["任务"]["B1"].value, "reference_image_urls")
        self.assertEqual(template["任务"]["C1"].value, "reference_roles")
        self.assertEqual(template["任务"]["D1"].value, "product_box")
        self.assertEqual(template["任务"]["G1"].value, "quality")
        self.assertEqual(template["任务"]["G2"].value, "low")
        self.assertGreater(len(template["任务"].data_validations.dataValidation), 0)

    def test_03_manual_multiple_videos(self) -> None:
        videos = [
            UploadFile(filename="one.mp4", file=io.BytesIO(b"video-one")),
            UploadFile(filename="two.mov", file=io.BytesIO(b"video-two")),
        ]
        references = [
            UploadFile(filename="person.jpg", file=io.BytesIO(b"person")),
            UploadFile(filename="product.png", file=io.BytesIO(b"product")),
        ]
        result = asyncio.run(self.main.create_manual_task(
            upload_videos=videos,
            reference_images=references,
            prompt="Picture 1 is the person. Picture 2 is the product.",
            aspect_ratio="16:9",
            seed=None,
            video_durations="[5.2, 8.4]",
        ))
        self.assertEqual(result["created"], 2)
        with self.main.connect_db() as conn:
            rows = conn.execute(
                "SELECT name,duration,reference_images,quality FROM tasks WHERE id IN (?,?) ORDER BY name",
                result["task_ids"],
            ).fetchall()
        self.assertEqual([row["name"] for row in rows], ["one", "two"])
        self.assertEqual([row["duration"] for row in rows], [5.2, 8.4])
        self.assertTrue(all(len(json.loads(row["reference_images"])) == 2 for row in rows))
        self.assertTrue(all(row["quality"] == "low" for row in rows))

    def test_04_gpu_status_parser(self) -> None:
        parsed = self.main.parse_nvidia_smi(
            "0, NVIDIA L40S, 87, 36120, 46068, 68, 287.4, 350.0\n"
        )
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["name"], "NVIDIA L40S")
        self.assertEqual(parsed[0]["utilization"], 87.0)
        self.assertEqual(parsed[0]["memory_used_mb"], 36120.0)
        self.assertAlmostEqual(parsed[0]["memory_percent"], 78.4)

    def test_05_comfy_sampling_progress(self) -> None:
        workflow = {"125": {"class_type": "SamplerCustomAdvanced", "_meta": {"title": "采样器"}}}
        with patch.object(self.main, "mark_task") as mark:
            finished = self.main.apply_comfy_event(
                "task-1", "prompt-1", workflow,
                {"type": "progress", "data": {
                    "prompt_id": "prompt-1", "node": "125", "value": 10, "max": 20,
                }},
            )
        self.assertFalse(finished)
        values = mark.call_args.kwargs
        self.assertEqual(values["stage"], "采样生成视频（10/20）")
        self.assertGreater(values["progress"], 50)

    def test_06_long_video_segment_plan(self) -> None:
        self.assertEqual(self.main.segment_durations(15), [15.0])
        self.assertEqual(self.main.segment_durations(16), [8.0, 8.0])
        parts = self.main.segment_durations(31)
        self.assertEqual(len(parts), 3)
        self.assertAlmostEqual(sum(parts), 31.0, places=3)
        self.assertTrue(all(4 <= value <= 15 for value in parts))

    def test_07_long_video_processes_parts_and_merges(self) -> None:
        folder = Path(self.temp.name) / "long-video"
        folder.mkdir(exist_ok=True)
        source = folder / "source.mp4"
        source.write_bytes(b"long-video")
        task_id = self.main.insert_task(
            prompt="Replace the person.", source_video=source,
            reference_images=[], duration=31.0,
            aspect_ratio="auto", seed=7, display_name="long.mp4",
        )
        with self.main.connect_db() as conn:
            row = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()

        observed: list[tuple[int, int, float]] = []

        def fake_split(_, durations, work_folder):
            parts = []
            for index, _duration in enumerate(durations, start=1):
                part = work_folder / f"source_part_{index:03d}.mp4"
                part.write_bytes(b"source-part")
                parts.append(part)
            return parts

        def fake_segment(parent_id, segment_row, index, count, destination):
            self.assertEqual(parent_id, task_id)
            observed.append((index, count, segment_row["duration"]))
            destination.write_bytes(b"generated-part")
            return True

        def fake_concat(parts, destination, _work_folder):
            self.assertEqual(len(parts), 3)
            destination.write_bytes(b"merged-video")

        with patch.object(self.main, "split_source_video", side_effect=fake_split), patch.object(
            self.main, "run_comfy_segment", side_effect=fake_segment
        ), patch.object(self.main, "concat_generated_videos", side_effect=fake_concat):
            self.main.process_task(row)

        self.assertEqual(len(observed), 3)
        self.assertAlmostEqual(sum(item[2] for item in observed), 31.0, places=3)
        with self.main.connect_db() as conn:
            completed = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["segment_count"], 3)
        self.assertEqual(completed["current_segment"], 3)
        self.assertTrue(Path(completed["output_path"]).is_file())

    def test_08_precision_pipeline_orders_product_face_and_audio(self) -> None:
        folder = Path(self.temp.name) / "precision"
        folder.mkdir(exist_ok=True)
        source = folder / "source.mp4"
        face = folder / "face.jpg"
        product = folder / "product.png"
        source.write_bytes(b"video")
        face.write_bytes(b"face")
        product.write_bytes(b"product")
        task_id = self.main.insert_task(
            prompt="只替换商品和脸", source_video=source,
            reference_images=[face, product], reference_roles=["face", "product"],
            product_box=[0.2, 0.3, 0.6, 0.8], precision_mode=True,
            duration=6.0, aspect_ratio="auto", seed=12, display_name="precise.mp4",
        )
        with self.main.connect_db() as conn:
            row = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()

        calls: list[str] = []

        def fake_segment(_parent, segment_row, _index, _count, destination):
            self.assertEqual(self.main.row_reference_paths(segment_row), [str(product)])
            destination.write_bytes(b"candidate")
            calls.append("h3")
            return True

        def fake_composite(_source, _candidate, box, destination, _work):
            self.assertEqual(box, [0.2, 0.3, 0.6, 0.8])
            destination.write_bytes(b"product-composite")
            calls.append("sam2")

        def fake_facefusion(refs, _target, destination, _work):
            self.assertEqual(len(refs), 1)
            destination.write_bytes(b"face-output")
            calls.append("facefusion")

        def fake_audio(_processed, _source, destination):
            destination.write_bytes(b"final")
            calls.append("audio")

        with patch.object(self.main, "probe_video_dimensions", return_value=(1920, 1080)), patch.object(
            self.main, "run_comfy_segment", side_effect=fake_segment
        ), patch.object(self.main, "run_product_mask_composite", side_effect=fake_composite), patch.object(
            self.main, "run_facefusion", side_effect=fake_facefusion
        ), patch.object(self.main, "restore_original_audio", side_effect=fake_audio):
            self.main.process_task(row)

        self.assertEqual(calls, ["h3", "sam2", "facefusion", "audio"])
        with self.main.connect_db() as conn:
            completed = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["stage"], "精准替换完成")
        self.assertEqual((completed["output_width"], completed["output_height"]), (1920, 1080))


if __name__ == "__main__":
    unittest.main()
