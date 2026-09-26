from __future__ import annotations

import io
import json
import math
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import threading
import time
import urllib.error
import urllib.request
import uuid
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import urlencode, urlparse

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parent.parent
STATIC_DIR = ROOT / "static"
DATA_DIR = Path(os.getenv("H3_STUDIO_DATA_DIR", ROOT / "data")).resolve()
UPLOAD_DIR = DATA_DIR / "uploads"
OUTPUT_DIR = DATA_DIR / "outputs"
WORK_DIR = DATA_DIR / "work"
DB_PATH = DATA_DIR / "studio.db"
COMFYUI_URL = os.getenv("H3_COMFYUI_URL", "http://127.0.0.1:6008").rstrip("/")
COMFYUI_DIR = Path(os.getenv("H3_COMFYUI_DIR", "/root/autodl-tmp/ComfyUI")).resolve()
COMFYUI_INPUT_DIR = Path(os.getenv("H3_COMFYUI_INPUT_DIR", COMFYUI_DIR / "input")).resolve()
COMFYUI_OUTPUT_DIR = Path(os.getenv("H3_COMFYUI_OUTPUT_DIR", COMFYUI_DIR / "output")).resolve()
COMFYUI_WORKFLOW = Path(
    os.getenv("H3_COMFYUI_WORKFLOW", ROOT / "workflows" / "minimax_h3_ref2va_api.json")
).resolve()
COMFYUI_TIMEOUT = max(300, int(os.getenv("H3_COMFYUI_TIMEOUT", "10800")))
COMFYUI_UNET = os.getenv(
    "H3_COMFYUI_UNET", "minimax_h3_ref2va_pruned_int8_convrot.safetensors"
)
COMFYUI_CLIP = os.getenv(
    "H3_COMFYUI_CLIP", "qwen3vl_32b_minimax_h3_nvfp4_awq.safetensors"
)
COMFYUI_VIDEO_VAE = os.getenv(
    "H3_COMFYUI_VIDEO_VAE", "minimax_h3_video_vae_int8_convrot.safetensors"
)
COMFYUI_AUDIO_VAE = os.getenv("H3_COMFYUI_AUDIO_VAE", "minimax_h3_audio_vae_fp32.safetensors")
POLL_SECONDS = max(1.0, float(os.getenv("H3_POLL_SECONDS", "3")))
MAX_UPLOAD_GB = max(1, int(os.getenv("H3_MAX_UPLOAD_GB", "20")))
MAX_UPLOAD_BYTES = MAX_UPLOAD_GB * 1024**3
MAX_SEGMENT_SECONDS = 15.0
MIN_VIDEO_SECONDS = 4.0
QUALITY_LEVELS = {"low", "standard", "high"}

VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
worker_stop = threading.Event()
worker_thread: threading.Thread | None = None
gpu_cache_lock = threading.Lock()
gpu_cache_at = 0.0
gpu_cache_value: dict[str, Any] = {"available": False, "gpus": []}


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


@contextmanager
def connect_db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db() -> None:
    for folder in (DATA_DIR, UPLOAD_DIR, OUTPUT_DIR, WORK_DIR):
        folder.mkdir(parents=True, exist_ok=True)
    with connect_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS tasks (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                status TEXT NOT NULL,
                prompt TEXT NOT NULL,
                source_video TEXT NOT NULL,
                character_image TEXT,
                product_image TEXT,
                reference_images TEXT,
                duration REAL NOT NULL,
                aspect_ratio TEXT NOT NULL,
                source_start REAL NOT NULL DEFAULT 0,
                seed INTEGER NOT NULL,
                engine_job_id TEXT,
                output_path TEXT,
                error TEXT,
                progress INTEGER NOT NULL DEFAULT 0,
                stage TEXT,
                cancel_requested INTEGER NOT NULL DEFAULT 0,
                deleted_at TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                started_at TEXT,
                finished_at TEXT,
                segment_count INTEGER NOT NULL DEFAULT 1,
                current_segment INTEGER NOT NULL DEFAULT 0,
                quality TEXT NOT NULL DEFAULT 'low',
                output_width INTEGER,
                output_height INTEGER
            );
            CREATE INDEX IF NOT EXISTS idx_tasks_queue
                ON tasks(status, deleted_at, created_at);
            """
        )
        columns = {row["name"] for row in conn.execute("PRAGMA table_info(tasks)").fetchall()}
        if "reference_images" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN reference_images TEXT")
        if "stage" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN stage TEXT")
        if "segment_count" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN segment_count INTEGER NOT NULL DEFAULT 1")
        if "current_segment" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN current_segment INTEGER NOT NULL DEFAULT 0")
        if "quality" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN quality TEXT NOT NULL DEFAULT 'low'")
            conn.execute("UPDATE tasks SET quality='high'")
        if "output_width" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN output_width INTEGER")
        if "output_height" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN output_height INTEGER")
        conn.execute(
            "UPDATE tasks SET status='queued', progress=0, stage='等待队列', started_at=NULL, "
            "updated_at=? WHERE status IN ('starting','running') AND deleted_at IS NULL",
            (utcnow(),),
        )


def validate_extension(filename: str, allowed: set[str], label: str) -> None:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in allowed:
        expected = "、".join(sorted(allowed))
        raise HTTPException(400, f"{label}格式不支持，请使用 {expected}")


def segment_durations(duration: float) -> list[float]:
    """Split a source duration into nearly equal H3-compatible 4–15 second parts."""
    if not math.isfinite(duration) or duration < MIN_VIDEO_SECONDS:
        raise ValueError(f"上传视频时长为 {duration:.2f} 秒；视频不能短于 {MIN_VIDEO_SECONDS:g} 秒")
    count = max(1, math.ceil(duration / MAX_SEGMENT_SECONDS))
    segment = duration / count
    if segment < MIN_VIDEO_SECONDS:
        raise ValueError("无法把视频切分为符合 MiniMax-H3 要求的片段")
    parts = [round(segment, 3) for _ in range(count)]
    parts[-1] = round(duration - sum(parts[:-1]), 3)
    return parts


def validate_task_values(duration: float, aspect_ratio: str, quality: str = "low") -> None:
    segment_durations(duration)
    if aspect_ratio not in {"auto", "16:9", "9:16", "1:1", "4:3", "3:4", "21:9"}:
        raise ValueError("不支持的画面比例")
    if quality not in QUALITY_LEVELS:
        raise ValueError("不支持的视频清晰度")


def is_http_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def validate_public_media_url(value: str, label: str) -> str:
    value = value.strip()
    if not is_http_url(value):
        raise ValueError(f"{label}必须是完整的 http:// 或 https:// URL")
    if urlparse(value).username or urlparse(value).password:
        raise ValueError(f"{label}不能在 URL 中包含用户名或密码")
    return value


def media_display_name(value: str | Path) -> str:
    text = str(value)
    if is_http_url(text):
        return Path(urlparse(text).path).name or urlparse(text).netloc
    return Path(text).name


def probe_video_duration(source: str | Path, hint: float | None = None) -> float:
    duration: float | None = None
    try:
        completed = subprocess.run(
            [
                "ffprobe", "-v", "error", "-show_entries", "format=duration",
                "-of", "default=noprint_wrappers=1:nokey=1", str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        duration = float(completed.stdout.strip())
    except (FileNotFoundError, subprocess.SubprocessError, ValueError):
        if hint is not None and hint > 0:
            duration = float(hint)
    if duration is None:
        raise ValueError(f"无法读取视频“{media_display_name(source)}”的时长，请确认 URL 可公开访问或视频可正常播放")
    duration = round(duration, 3)
    validate_task_values(duration, "16:9")
    return duration


async def save_upload(upload: UploadFile, destination: Path, allowed: set[str], label: str) -> Path:
    validate_extension(upload.filename or "", allowed, label)
    destination.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    try:
        with destination.open("wb") as target:
            while chunk := await upload.read(1024 * 1024):
                written += len(chunk)
                if written > MAX_UPLOAD_BYTES:
                    raise HTTPException(413, f"单个文件不能超过 {MAX_UPLOAD_GB} GB")
                target.write(chunk)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        await upload.close()
    return destination.resolve()


def safe_name(filename: str) -> str:
    raw = Path(filename or "file").name
    cleaned = "".join(c if c.isalnum() or c in "._-" else "_" for c in raw)
    return cleaned[:160] or "file"


def task_to_dict(row: sqlite3.Row, queue_position: int | None = None) -> dict[str, Any]:
    result = {key: row[key] for key in row.keys()}
    for field in ("source_video", "character_image", "product_image", "output_path"):
        if result.get(field):
            result[field] = media_display_name(result[field])
    result["reference_images"] = [media_display_name(path) for path in row_reference_paths(row)]
    result["cancel_requested"] = bool(result["cancel_requested"])
    result["queue_position"] = queue_position
    result["download_url"] = f"/api/tasks/{row['id']}/output" if row["status"] == "completed" else None
    return result


def insert_task(
    *, prompt: str, source_video: str | Path, reference_images: list[str | Path], duration: float,
    aspect_ratio: str, seed: int | None, quality: str = "low", display_name: str | None = None,
) -> str:
    validate_task_values(duration, aspect_ratio, quality)
    source_text = str(source_video)
    if is_http_url(source_text):
        validate_public_media_url(source_text, "upload_video_url")
    else:
        source_path = Path(source_text)
        if not source_path.is_file() or source_path.suffix.lower() not in VIDEO_EXTENSIONS:
            raise ValueError("upload_video 必须是已上传的受支持视频文件")
    if len(reference_images) > 9:
        raise ValueError("参考图片最多上传 9 张")
    for image in reference_images:
        image_text = str(image)
        if is_http_url(image_text):
            validate_public_media_url(image_text, "reference_image_url")
        else:
            image_path = Path(image_text)
            if not image_path.is_file() or image_path.suffix.lower() not in IMAGE_EXTENSIONS:
                raise ValueError("reference_images 必须是已上传的受支持图片文件")
    task_id = uuid.uuid4().hex
    now = utcnow()
    actual_seed = seed if seed is not None else secrets.randbelow(2_147_483_647)
    name = Path(display_name or media_display_name(source_video)).stem
    part_count = len(segment_durations(duration))
    with connect_db() as conn:
        conn.execute(
            """INSERT INTO tasks (
                id,name,status,prompt,source_video,reference_images,
                duration,aspect_ratio,source_start,seed,stage,segment_count,current_segment,quality,
                created_at,updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                task_id, name or f"视频-{task_id[:6]}", "queued", prompt.strip(),
                source_text, json.dumps([str(path) for path in reference_images], ensure_ascii=False),
                duration, aspect_ratio, 0, actual_seed, "等待队列", part_count, 0, quality, now, now,
            ),
        )
    return task_id


def row_reference_paths(row: sqlite3.Row | dict[str, Any]) -> list[str]:
    keys = set(row.keys())
    raw = row["reference_images"] if "reference_images" in keys else None
    if raw:
        try:
            parsed = json.loads(raw)
            if isinstance(parsed, list):
                return [str(path) for path in parsed if path]
        except (TypeError, json.JSONDecodeError):
            pass
    legacy = []
    for field in ("character_image", "product_image"):
        if field in keys and row[field]:
            legacy.append(str(row[field]))
    return legacy


def build_h3_prompt(row: sqlite3.Row | dict[str, Any]) -> str:
    definitions = [
        "<Video 1> is the source video. Preserve its shot order, camera motion, subject motion, "
        "timing, composition, lighting, background, and synchronized soundtrack unless explicitly changed below."
    ]
    reference_paths = row_reference_paths(row)
    instructions: list[str] = []
    for picture_index, _ in enumerate(reference_paths, start=1):
        definitions.append(
            f"<Picture {picture_index}> is a user-provided visual reference for the replacement person, product, or both."
        )
    if reference_paths:
        instructions.append(
            "Use the supplied <Picture> references according to the user's request. For a person, preserve identity, face, hair, "
            "pose, motion, gaze, scale, placement, and timing. For a product, preserve category, shape, colors, branding, packaging, "
            "perspective, occlusion, hand contact, reflections, shadows, and temporal consistency."
        )
    if not instructions:
        instructions.append("Apply the requested edit to <Video 1> while preserving all unspecified content.")

    user_request = row["prompt"].strip() or "Keep the result natural, photorealistic, and temporally consistent."
    return (
        "subject_definitions:\n" + "\n".join(definitions) +
        "\n\nsummary:\n[video editing + identity/product reference + audio reuse] "
        "The target is a faithful edited version of <Video 1>. " + " ".join(instructions) +
        " Preserve the original environment and synchronized audio. Do not add extra people, products, text, logos, cuts, or camera moves."
        "\n\nretention_analysis:\n<Video 1>: fully_preserved except for the explicitly requested replacements. "
        "Motion, timing, framing, background, lighting, and audio remain consistent."
        "\n\ndetailed_description:\n" + user_request
    )


def http_json(method: str, url: str, payload: dict[str, Any] | None = None, timeout: int = 60) -> dict[str, Any]:
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(url, data=body, method=method)
    req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            content = response.read()
            return json.loads(content.decode("utf-8")) if content else {}
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:2000]
        raise RuntimeError(f"ComfyUI 返回 HTTP {exc.code}: {detail}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"无法连接 ComfyUI {COMFYUI_URL}: {exc.reason}") from exc


def parse_nvidia_smi(output: str) -> list[dict[str, Any]]:
    gpus: list[dict[str, Any]] = []
    for line in output.splitlines():
        values = [value.strip() for value in line.split(",")]
        if len(values) != 8:
            continue
        try:
            index, name, utilization, memory_used, memory_total, temperature, power_draw, power_limit = values
            used = float(memory_used)
            total = float(memory_total)
            gpus.append({
                "index": int(index),
                "name": name,
                "utilization": float(utilization),
                "memory_used_mb": used,
                "memory_total_mb": total,
                "memory_percent": round(used / total * 100, 1) if total else 0,
                "temperature_c": float(temperature),
                "power_draw_w": float(power_draw),
                "power_limit_w": float(power_limit),
            })
        except ValueError:
            continue
    return gpus


def gpu_status() -> dict[str, Any]:
    global gpu_cache_at, gpu_cache_value
    now = time.monotonic()
    with gpu_cache_lock:
        if now - gpu_cache_at < 2:
            return gpu_cache_value
        try:
            completed = subprocess.run(
                [
                    "nvidia-smi",
                    "--query-gpu=index,name,utilization.gpu,memory.used,memory.total,"
                    "temperature.gpu,power.draw,power.limit",
                    "--format=csv,noheader,nounits",
                ],
                check=True, capture_output=True, text=True, timeout=3,
            )
            gpus = parse_nvidia_smi(completed.stdout)
            gpu_cache_value = {"available": bool(gpus), "gpus": gpus, "error": None}
        except (FileNotFoundError, subprocess.SubprocessError) as exc:
            gpu_cache_value = {"available": False, "gpus": [], "error": type(exc).__name__}
        gpu_cache_at = now
        return gpu_cache_value


def comfy_websocket_url(client_id: str) -> str:
    parsed = urlparse(COMFYUI_URL)
    scheme = "wss" if parsed.scheme == "https" else "ws"
    return f"{scheme}://{parsed.netloc}/ws?{urlencode({'clientId': client_id})}"


def open_comfy_websocket(client_id: str):
    try:
        from websockets.sync.client import connect
        return connect(comfy_websocket_url(client_id), open_timeout=3, close_timeout=1)
    except Exception:
        return None


def comfy_node_stage(workflow: dict[str, Any], node_id: Any) -> tuple[str, int]:
    node = workflow.get(str(node_id), {})
    class_type = str(node.get("class_type") or "")
    title = str((node.get("_meta") or {}).get("title") or class_type)
    if class_type in {"LoadVideo", "GetVideoComponents", "LoadImage"}:
        return "读取参考素材", 18
    if class_type in {"UNETLoader", "CLIPLoader", "VAELoader"}:
        return f"加载模型：{title}", 22
    if class_type == "MiniMaxH3ReferenceToVideo":
        return "编码视频、图片与提示词", 30
    if class_type in {"BasicScheduler", "BasicGuider", "KSamplerSelect", "RandomNoise"}:
        return "准备采样参数", 34
    if class_type == "SamplerCustomAdvanced":
        return "采样生成视频", 36
    if class_type == "VAEDecode":
        return "解码视频画面", 90
    if class_type == "VAEDecodeAudio":
        return "解码音频", 92
    if class_type == "CreateVideo":
        return "合成视频与音频", 95
    if class_type == "SaveVideo":
        return "保存生成视频", 97
    return (f"执行节点：{title}" if title else "ComfyUI 生成中"), 20


def segment_stage(stage: str, segment_index: int, segment_count: int) -> str:
    if segment_count <= 1:
        return stage
    return f"第 {segment_index}/{segment_count} 段 · {stage}"


def segment_progress(local_progress: int | float, segment_index: int, segment_count: int) -> int:
    local = min(100.0, max(0.0, float(local_progress)))
    count = max(1, segment_count)
    overall = 8 + ((segment_index - 1 + local / 100) / count) * 86
    return min(94, max(8, round(overall)))


def mark_segment_task(
    task_id: str, segment_index: int, segment_count: int, stage: str, local_progress: int | float,
    **values: Any,
) -> None:
    mark_task(
        task_id,
        current_segment=segment_index,
        segment_count=segment_count,
        stage=segment_stage(stage, segment_index, segment_count),
        progress=segment_progress(local_progress, segment_index, segment_count),
        **values,
    )


def apply_comfy_event(
    task_id: str, prompt_id: str, workflow: dict[str, Any], message: dict[str, Any],
    segment_index: int = 1, segment_count: int = 1,
) -> bool:
    event_type = message.get("type")
    data = message.get("data") or {}
    if data.get("prompt_id") not in {None, prompt_id}:
        return False
    if event_type == "execution_start":
        mark_segment_task(task_id, segment_index, segment_count, "ComfyUI 开始执行", 17)
    elif event_type == "executing":
        node_id = data.get("node")
        if node_id is None and data.get("prompt_id") == prompt_id:
            mark_segment_task(task_id, segment_index, segment_count, "ComfyUI 执行完成", 94)
            return True
        stage, progress = comfy_node_stage(workflow, node_id)
        mark_segment_task(task_id, segment_index, segment_count, stage, progress)
    elif event_type == "progress":
        value = float(data.get("value") or 0)
        maximum = float(data.get("max") or 0)
        stage, base = comfy_node_stage(workflow, data.get("node"))
        if maximum > 0:
            percent = min(1.0, max(0.0, value / maximum))
            progress = max(base, min(89, 36 + round(percent * 53)))
            stage = f"{stage}（{value:g}/{maximum:g}）"
            mark_segment_task(task_id, segment_index, segment_count, stage, progress)
    elif event_type == "progress_state":
        nodes = data.get("nodes") or {}
        running = next((item for item in nodes.values() if item.get("state") == "running"), None)
        if running:
            display_id = running.get("display_node_id") or running.get("real_node_id") or running.get("node_id")
            synthetic = {"type": "progress", "data": {
                "prompt_id": prompt_id, "node": display_id,
                "value": running.get("value"), "max": running.get("max"),
            }}
            apply_comfy_event(
                task_id, prompt_id, workflow, synthetic, segment_index, segment_count
            )
    return False


def copy_media_to_comfy(source: str | Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    source_text = str(source)
    written = 0
    try:
        if is_http_url(source_text):
            request = urllib.request.Request(source_text, headers={"User-Agent": "H3-Studio/1.0"})
            with urllib.request.urlopen(request, timeout=120) as response, destination.open("wb") as target:
                while chunk := response.read(1024 * 1024):
                    written += len(chunk)
                    if written > MAX_UPLOAD_BYTES:
                        raise RuntimeError(f"远程素材不能超过 {MAX_UPLOAD_GB} GB")
                    target.write(chunk)
        else:
            shutil.copy2(Path(source_text), destination)
    except Exception:
        destination.unlink(missing_ok=True)
        raise


def prepare_comfy_inputs(row: sqlite3.Row) -> tuple[str, list[str], Path]:
    task_folder = COMFYUI_INPUT_DIR / "h3_studio" / row["id"]
    task_folder.mkdir(parents=True, exist_ok=True)
    try:
        source_suffix = Path(urlparse(str(row["source_video"])).path).suffix.lower() or ".mp4"
        if source_suffix not in VIDEO_EXTENSIONS:
            source_suffix = ".mp4"
        source_path = task_folder / f"source{source_suffix}"
        copy_media_to_comfy(row["source_video"], source_path)

        reference_names: list[str] = []
        for index, source in enumerate(row_reference_paths(row), start=1):
            suffix = Path(urlparse(str(source)).path).suffix.lower() or ".jpg"
            if suffix not in IMAGE_EXTENSIONS:
                suffix = ".jpg"
            destination = task_folder / f"reference_{index:02d}{suffix}"
            copy_media_to_comfy(source, destination)
            reference_names.append(destination.relative_to(COMFYUI_INPUT_DIR).as_posix())
        return source_path.relative_to(COMFYUI_INPUT_DIR).as_posix(), reference_names, task_folder
    except Exception:
        shutil.rmtree(task_folder, ignore_errors=True)
        raise


QUALITY_DIMENSIONS = {
    "low": {
        "16:9": (672, 384),
        "9:16": (384, 672),
        "1:1": (512, 512),
        "4:3": (576, 448),
        "3:4": (448, 576),
        "21:9": (768, 320),
    },
    "standard": {
        "16:9": (832, 480),
        "9:16": (480, 832),
        "1:1": (640, 640),
        "4:3": (736, 544),
        "3:4": (544, 736),
        "21:9": (960, 416),
    },
    "high": {
        "16:9": (1344, 768),
        "9:16": (768, 1344),
        "1:1": (1024, 1024),
        "4:3": (1024, 768),
        "3:4": (768, 1024),
        "21:9": (1568, 672),
    },
}


def source_aspect_ratio(source: str | Path) -> str:
    try:
        completed = subprocess.run(
            [
                "ffprobe", "-v", "error", "-select_streams", "v:0",
                "-show_entries", "stream=width,height", "-of", "csv=p=0:s=x", str(source),
            ],
            check=True, capture_output=True, text=True, timeout=30,
        )
        width, height = (int(value) for value in completed.stdout.strip().split("x", 1))
        ratio = width / height
        dimensions = QUALITY_DIMENSIONS["high"]
        return min(dimensions, key=lambda name: abs(ratio - dimensions[name][0] / dimensions[name][1]))
    except (FileNotFoundError, subprocess.SubprocessError, ValueError, ZeroDivisionError):
        return "16:9"


def next_workflow_node_id(workflow: dict[str, Any]) -> str:
    return str(max(int(node_id) for node_id in workflow) + 1)


def build_comfy_workflow(row: sqlite3.Row, source_upload: str, reference_uploads: list[str]) -> dict[str, Any]:
    if not COMFYUI_WORKFLOW.is_file():
        raise RuntimeError(f"找不到 ComfyUI 工作流：{COMFYUI_WORKFLOW}")
    workflow = json.loads(COMFYUI_WORKFLOW.read_text(encoding="utf-8"))
    sampler = workflow["136"]["inputs"]
    local_source = COMFYUI_INPUT_DIR / source_upload
    ratio = row["aspect_ratio"] if row["aspect_ratio"] != "auto" else source_aspect_ratio(local_source)
    quality = str(row.get("quality") or "low") if isinstance(row, dict) else str(row["quality"] or "low")
    if quality not in QUALITY_LEVELS:
        quality = "low"
    width, height = QUALITY_DIMENSIONS[quality][ratio]
    sampler["width"] = width
    sampler["height"] = height
    workflow["132"]["inputs"]["value"] = float(row["duration"])
    workflow["129"]["inputs"]["noise_seed"] = int(row["seed"])
    workflow["138"]["inputs"]["value"] = build_h3_prompt(row)
    workflow["92"]["inputs"]["filename_prefix"] = f"h3_studio/{row['id']}"
    workflow["127"]["inputs"]["unet_name"] = COMFYUI_UNET
    workflow["128"]["inputs"]["clip_name"] = COMFYUI_CLIP
    workflow["119"]["inputs"]["vae_name"] = COMFYUI_VIDEO_VAE
    workflow["120"]["inputs"]["vae_name"] = COMFYUI_AUDIO_VAE

    for key in list(sampler):
        if key.startswith(("ref_images.ref_image_", "ref_videos.ref_video_", "ref_video_audios.ref_video_audio_")):
            del sampler[key]
    for node_id in list(workflow):
        if workflow[node_id]["class_type"] in {"LoadImage", "LoadVideo", "GetVideoComponents"}:
            del workflow[node_id]

    load_video_id = next_workflow_node_id(workflow)
    workflow[load_video_id] = {
        "class_type": "LoadVideo", "inputs": {"file": source_upload},
        "_meta": {"title": "上传视频"},
    }
    components_id = next_workflow_node_id(workflow)
    workflow[components_id] = {
        "class_type": "GetVideoComponents", "inputs": {"video": [load_video_id, 0]},
        "_meta": {"title": "读取视频画面与声音"},
    }
    sampler["ref_videos.ref_video_0"] = [components_id, 0]
    sampler["ref_video_audios.ref_video_audio_0"] = [components_id, 1]

    for index, uploaded_name in enumerate(reference_uploads[:9]):
        node_id = next_workflow_node_id(workflow)
        workflow[node_id] = {
            "class_type": "LoadImage", "inputs": {"image": uploaded_name},
            "_meta": {"title": f"参考图片 {index + 1}"},
        }
        sampler[f"ref_images.ref_image_{index}"] = [node_id, 0]
    return workflow


def try_cancel_comfy(prompt_id: str | None) -> None:
    if not prompt_id:
        return
    try:
        http_json("POST", f"{COMFYUI_URL}/queue", {"delete": [prompt_id]}, timeout=10)
    except Exception:
        pass
    try:
        http_json("POST", f"{COMFYUI_URL}/interrupt", {}, timeout=10)
    except Exception:
        pass


def comfy_history_output(record: dict[str, Any]) -> dict[str, str]:
    status = record.get("status") or {}
    if status.get("status_str") == "error":
        messages = [item[1] for item in status.get("messages", []) if item and item[0] == "execution_error"]
        raise RuntimeError(f"ComfyUI 生成失败：{messages[-1] if messages else status}")
    node_output = (record.get("outputs") or {}).get("92") or {}
    for field in ("images", "videos", "gifs"):
        entries = node_output.get(field) or []
        if entries:
            return entries[0]
    raise RuntimeError("ComfyUI 已结束，但没有返回 SaveVideo 输出")


def download_comfy_output(output: dict[str, str], destination: Path) -> None:
    filename = Path(output.get("filename", "")).name
    subfolder = Path(output.get("subfolder", ""))
    source = (COMFYUI_OUTPUT_DIR / subfolder / filename).resolve()
    output_root = COMFYUI_OUTPUT_DIR.resolve()
    if source.is_file() and (source == output_root or output_root in source.parents):
        shutil.copy2(source, destination)
        source.unlink(missing_ok=True)
        return
    query = urlencode({
        "filename": filename, "subfolder": output.get("subfolder", ""),
        "type": output.get("type", "output"),
    })
    req = urllib.request.Request(f"{COMFYUI_URL}/view?{query}")
    try:
        with urllib.request.urlopen(req, timeout=900) as response, destination.open("wb") as target:
            shutil.copyfileobj(response, target, length=1024 * 1024)
    except Exception:
        destination.unlink(missing_ok=True)
        raise


def run_ffmpeg(arguments: list[str], purpose: str) -> None:
    try:
        completed = subprocess.run(
            ["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", *arguments],
            capture_output=True,
            text=True,
            timeout=COMFYUI_TIMEOUT,
        )
    except FileNotFoundError as exc:
        raise RuntimeError("找不到 ffmpeg，请重新运行安装脚本") from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"{purpose}超时") from exc
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "未知错误").strip()[-2000:]
        raise RuntimeError(f"{purpose}失败：{detail}")


def split_source_video(source: Path, durations: list[float], work_folder: Path) -> list[Path]:
    work_folder.mkdir(parents=True, exist_ok=True)
    parts: list[Path] = []
    start = 0.0
    for index, duration in enumerate(durations, start=1):
        destination = work_folder / f"source_part_{index:03d}.mp4"
        run_ffmpeg(
            [
                "-ss", f"{start:.3f}", "-i", str(source), "-t", f"{duration:.3f}",
                "-map", "0:v:0", "-map", "0:a?", "-c:v", "libx264", "-preset", "veryfast",
                "-crf", "18", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k",
                "-movflags", "+faststart", str(destination),
            ],
            f"切割第 {index}/{len(durations)} 段视频",
        )
        if not destination.is_file() or destination.stat().st_size == 0:
            raise RuntimeError(f"切割第 {index}/{len(durations)} 段视频后没有生成文件")
        parts.append(destination)
        start += duration
    return parts


def concat_generated_videos(parts: list[Path], destination: Path, work_folder: Path) -> None:
    if not parts:
        raise RuntimeError("没有可合并的生成片段")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if len(parts) == 1:
        shutil.move(str(parts[0]), destination)
        return
    concat_file = work_folder / "generated_parts.txt"
    concat_lines = []
    for path in parts:
        escaped_path = str(path.resolve()).replace("'", "'\\''")
        concat_lines.append(f"file '{escaped_path}'\n")
    concat_file.write_text("".join(concat_lines), encoding="utf-8")
    copy_error: Exception | None = None
    try:
        run_ffmpeg(
            ["-f", "concat", "-safe", "0", "-i", str(concat_file), "-c", "copy", "-movflags", "+faststart", str(destination)],
            "无损合并生成片段",
        )
    except Exception as exc:
        copy_error = exc
        destination.unlink(missing_ok=True)
        run_ffmpeg(
            [
                "-f", "concat", "-safe", "0", "-i", str(concat_file),
                "-c:v", "libx264", "-preset", "veryfast", "-crf", "18",
                "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart", str(destination),
            ],
            "重新编码并合并生成片段",
        )
    if not destination.is_file() or destination.stat().st_size == 0:
        raise RuntimeError(f"合并片段后没有生成文件；首次合并错误：{copy_error or '无'}")


def local_source_for_split(row: sqlite3.Row, work_folder: Path) -> Path:
    source = str(row["source_video"])
    if not is_http_url(source):
        return Path(source)
    suffix = Path(urlparse(source).path).suffix.lower()
    if suffix not in VIDEO_EXTENSIONS:
        suffix = ".mp4"
    destination = work_folder / f"downloaded_source{suffix}"
    copy_media_to_comfy(source, destination)
    return destination


def claim_next_task() -> sqlite3.Row | None:
    with connect_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute(
            "SELECT * FROM tasks WHERE status='queued' AND deleted_at IS NULL "
            "AND cancel_requested=0 ORDER BY created_at, id LIMIT 1"
        ).fetchone()
        if row is None:
            return None
        now = utcnow()
        conn.execute(
            "UPDATE tasks SET status='starting',progress=5,stage='准备任务素材',current_segment=0,started_at=?,updated_at=? WHERE id=?",
            (now, now, row["id"]),
        )
        return conn.execute("SELECT * FROM tasks WHERE id=?", (row["id"],)).fetchone()


def get_control_state(task_id: str) -> sqlite3.Row | None:
    with connect_db() as conn:
        return conn.execute(
            "SELECT cancel_requested,deleted_at FROM tasks WHERE id=?", (task_id,)
        ).fetchone()


def mark_task(task_id: str, **values: Any) -> None:
    if not values:
        return
    values["updated_at"] = utcnow()
    columns = ",".join(f"{key}=?" for key in values)
    with connect_db() as conn:
        conn.execute(f"UPDATE tasks SET {columns} WHERE id=?", (*values.values(), task_id))


def run_comfy_segment(
    task_id: str, segment_row: dict[str, Any], segment_index: int, segment_count: int,
    destination: Path,
) -> bool:
    prompt_id: str | None = None
    comfy_input_folder: Path | None = None
    websocket = None
    try:
        control = get_control_state(task_id)
        if control is None or control["cancel_requested"] or control["deleted_at"]:
            if control is not None and not control["deleted_at"]:
                mark_task(
                    task_id, status="cancelled", progress=0,
                    stage="任务已取消", finished_at=utcnow(),
                )
            return False
        mark_segment_task(task_id, segment_index, segment_count, "复制视频和参考图片", 8)
        source_upload, reference_uploads, comfy_input_folder = prepare_comfy_inputs(segment_row)
        mark_segment_task(task_id, segment_index, segment_count, "构建 MiniMax-H3 工作流", 12)
        workflow = build_comfy_workflow(segment_row, source_upload, reference_uploads)
        workflow_size = workflow["136"]["inputs"]
        mark_task(
            task_id,
            output_width=int(workflow_size["width"]),
            output_height=int(workflow_size["height"]),
        )
        client_id = f"h3-studio-{task_id}-{segment_index}"
        websocket = open_comfy_websocket(client_id)
        mark_segment_task(task_id, segment_index, segment_count, "提交任务到 ComfyUI", 14)
        submitted = http_json(
            "POST", f"{COMFYUI_URL}/prompt",
            {"prompt": workflow, "client_id": client_id}, timeout=120,
        )
        prompt_id = str(submitted.get("prompt_id") or "")
        if not prompt_id:
            raise RuntimeError(f"ComfyUI 未返回 prompt_id：{submitted}")
        mark_segment_task(
            task_id, segment_index, segment_count, "等待 ComfyUI 执行", 15,
            status="running", engine_job_id=prompt_id,
        )

        started = time.monotonic()
        history_record: dict[str, Any] | None = None
        while not worker_stop.is_set():
            control = get_control_state(task_id)
            if control is None or control["cancel_requested"] or control["deleted_at"]:
                try_cancel_comfy(prompt_id)
                if control is not None and not control["deleted_at"]:
                    mark_task(
                        task_id, status="cancelled", progress=0,
                        stage="任务已取消", finished_at=utcnow(),
                    )
                return False
            if websocket is not None:
                try:
                    raw_message = websocket.recv(timeout=POLL_SECONDS)
                    if isinstance(raw_message, str):
                        event_finished = apply_comfy_event(
                            task_id, prompt_id, workflow, json.loads(raw_message),
                            segment_index, segment_count,
                        )
                        if event_finished:
                            mark_segment_task(
                                task_id, segment_index, segment_count, "读取生成结果", 94
                            )
                except TimeoutError:
                    pass
                except Exception:
                    try:
                        websocket.close()
                    except Exception:
                        pass
                    websocket = None
            history = http_json("GET", f"{COMFYUI_URL}/history/{prompt_id}", timeout=30)
            history_record = history.get(prompt_id)
            if history_record is not None:
                break
            elapsed = time.monotonic() - started
            if elapsed > COMFYUI_TIMEOUT:
                try_cancel_comfy(prompt_id)
                raise TimeoutError(f"ComfyUI 生成超过 {COMFYUI_TIMEOUT} 秒")
            if websocket is None:
                local_progress = min(88, 18 + int(elapsed / 20))
                mark_segment_task(
                    task_id, segment_index, segment_count, "ComfyUI 生成中", local_progress
                )
                worker_stop.wait(POLL_SECONDS)
        if worker_stop.is_set():
            try_cancel_comfy(prompt_id)
            return False
        if history_record is None:
            raise RuntimeError("ComfyUI 未返回任务结果")

        mark_segment_task(task_id, segment_index, segment_count, "保存生成片段", 96)
        download_comfy_output(comfy_history_output(history_record), destination)
        try:
            http_json("POST", f"{COMFYUI_URL}/history", {"delete": [prompt_id]}, timeout=10)
        except Exception:
            pass
        return True
    finally:
        if websocket is not None:
            try:
                websocket.close()
            except Exception:
                pass
        if comfy_input_folder is not None:
            shutil.rmtree(comfy_input_folder, ignore_errors=True)


def process_task(row: sqlite3.Row) -> None:
    task_id = row["id"]
    work_folder = WORK_DIR / task_id
    destination = OUTPUT_DIR / f"{task_id}.mp4"
    try:
        durations = segment_durations(float(row["duration"]))
        segment_count = len(durations)
        mark_task(task_id, segment_count=segment_count, current_segment=0)
        work_folder.mkdir(parents=True, exist_ok=True)

        if segment_count > 1:
            mark_task(task_id, stage=f"正在切割源视频（共 {segment_count} 段）", progress=6)
            local_source = local_source_for_split(row, work_folder)
            source_parts = split_source_video(local_source, durations, work_folder)
        else:
            source_parts = [row["source_video"]]

        generated_parts: list[Path] = []
        for index, (source_part, part_duration) in enumerate(zip(source_parts, durations), start=1):
            segment_row = dict(row)
            segment_row["id"] = f"{task_id}_part_{index:03d}"
            segment_row["source_video"] = str(source_part)
            segment_row["duration"] = part_duration
            segment_destination = work_folder / f"generated_part_{index:03d}.mp4"
            if not run_comfy_segment(
                task_id, segment_row, index, segment_count, segment_destination
            ):
                return
            generated_parts.append(segment_destination)

        control = get_control_state(task_id)
        if control is None or control["cancel_requested"] or control["deleted_at"]:
            if control is not None and not control["deleted_at"]:
                mark_task(
                    task_id, status="cancelled", progress=0,
                    stage="任务已取消", finished_at=utcnow(),
                )
            return
        mark_task(task_id, progress=96, stage="合并生成片段" if segment_count > 1 else "保存生成视频")
        concat_generated_videos(generated_parts, destination, work_folder)
        control = get_control_state(task_id)
        if control is None or control["cancel_requested"] or control["deleted_at"]:
            destination.unlink(missing_ok=True)
            if control is not None and not control["deleted_at"]:
                mark_task(
                    task_id, status="cancelled", progress=0,
                    stage="任务已取消", finished_at=utcnow(),
                )
            return
        mark_task(
            task_id, status="completed", progress=100, output_path=str(destination),
            stage="生成完成", current_segment=segment_count, engine_job_id=None,
            finished_at=utcnow(), error=None,
        )
    except Exception as exc:
        destination.unlink(missing_ok=True)
        mark_task(
            task_id, status="failed", progress=0, stage="生成失败",
            error=str(exc)[:4000], finished_at=utcnow(),
        )
    finally:
        shutil.rmtree(work_folder, ignore_errors=True)


def worker_loop() -> None:
    while not worker_stop.is_set():
        row = claim_next_task()
        if row is None:
            worker_stop.wait(1.0)
            continue
        process_task(row)


@asynccontextmanager
async def lifespan(_: FastAPI):
    global worker_thread
    init_db()
    worker_stop.clear()
    worker_thread = threading.Thread(target=worker_loop, name="h3-task-worker", daemon=True)
    worker_thread.start()
    yield
    worker_stop.set()
    if worker_thread:
        worker_thread.join(timeout=5)


app = FastAPI(title="MiniMax H3 Video Studio", version="2.2.0", lifespan=lifespan)


@app.get("/api/config")
def config() -> dict[str, Any]:
    return {
        "max_upload_gb": MAX_UPLOAD_GB,
        "engine": "comfyui",
        "min_video_seconds": MIN_VIDEO_SECONDS,
        "max_segment_seconds": MAX_SEGMENT_SECONDS,
        "long_video_split": True,
        "default_quality": "low",
        "quality_dimensions": QUALITY_DIMENSIONS,
    }


@app.get("/api/system/status")
def system_status() -> dict[str, Any]:
    return {"gpu": gpu_status(), "updated_at": utcnow()}


@app.get("/health")
def health() -> dict[str, Any]:
    engine_ok = False
    try:
        req = urllib.request.Request(f"{COMFYUI_URL}/system_stats")
        with urllib.request.urlopen(req, timeout=2) as response:
            engine_ok = response.status < 500
    except Exception:
        pass
    return {
        "app": "ok", "engine": "ok" if engine_ok else "unavailable",
        "engine_name": "ComfyUI", "engine_url": COMFYUI_URL,
    }


@app.post("/api/tasks", status_code=201)
async def create_manual_task(
    upload_videos: Annotated[list[UploadFile], File(...)],
    reference_images: Annotated[list[UploadFile], File()] = [],
    prompt: Annotated[str, Form()] = "",
    aspect_ratio: Annotated[str, Form()] = "auto",
    quality: Annotated[str, Form()] = "low",
    seed: Annotated[int | None, Form()] = None,
    video_durations: Annotated[str, Form()] = "[]",
) -> dict[str, Any]:
    if not upload_videos or not any(item.filename for item in upload_videos):
        raise HTTPException(400, "请至少上传一个视频")
    if len([item for item in upload_videos if item.filename]) > 3:
        raise HTTPException(400, "一次最多上传 3 个视频")
    if len([item for item in reference_images if item.filename]) > 9:
        raise HTTPException(400, "参考图片最多上传 9 张")
    try:
        hints_raw = json.loads(video_durations)
        duration_hints = [float(value) for value in hints_raw] if isinstance(hints_raw, list) else []
    except (json.JSONDecodeError, TypeError, ValueError):
        duration_hints = []
    task_folder = UPLOAD_DIR / uuid.uuid4().hex
    created: list[str] = []
    try:
        saved_references: list[Path] = []
        for index, image in enumerate(reference_images, start=1):
            if image.filename:
                saved_references.append(await save_upload(
                    image, task_folder / f"ref_{index:02d}_{safe_name(image.filename)}", IMAGE_EXTENSIONS, "参考图片"
                ))

        prepared: list[tuple[Path, str, float]] = []
        for index, video in enumerate(upload_videos, start=1):
            if not video.filename:
                continue
            original_name = Path(video.filename).name
            saved_video = await save_upload(
                video, task_folder / f"video_{index:03d}_{safe_name(original_name)}", VIDEO_EXTENSIONS, "上传视频"
            )
            hint = duration_hints[index - 1] if index - 1 < len(duration_hints) else None
            duration = probe_video_duration(saved_video, hint)
            prepared.append((saved_video, original_name, duration))

        validate_task_values(prepared[0][2], aspect_ratio, quality)
        for saved_video, original_name, duration in prepared:
            created.append(insert_task(
                prompt=prompt, source_video=saved_video, reference_images=saved_references,
                duration=duration, aspect_ratio=aspect_ratio, seed=seed, quality=quality,
                display_name=original_name,
            ))
        return {"created": len(created), "task_ids": created, "status": "queued"}
    except HTTPException:
        if not created:
            shutil.rmtree(task_folder, ignore_errors=True)
        raise
    except Exception as exc:
        if not created:
            shutil.rmtree(task_folder, ignore_errors=True)
        raise HTTPException(400, str(exc)) from exc


def text_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def int_cell(value: Any, default: int | None = None) -> int | None:
    if value is None or str(value).strip() == "":
        return default
    return int(float(value))


@app.post("/api/tasks/import", status_code=201)
async def import_excel_tasks(
    excel: Annotated[UploadFile, File(...)],
) -> dict[str, Any]:
    validate_extension(excel.filename or "", EXCEL_EXTENSIONS, "Excel")
    created: list[str] = []
    errors: list[dict[str, Any]] = []
    try:
        excel_bytes = await excel.read()
        if len(excel_bytes) > 50 * 1024 * 1024:
            raise HTTPException(413, "Excel 文件不能超过 50 MB")
        await excel.close()

        workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        if "任务" not in workbook.sheetnames:
            raise HTTPException(400, "Excel 中缺少“任务”工作表，请使用页面下载的模板")
        sheet = workbook["任务"]
        rows = sheet.iter_rows(values_only=True)
        headers = [text_cell(value) for value in next(rows, [])]
        required = {"upload_video_url"}
        if not required.issubset(headers):
            raise HTTPException(400, "Excel 缺少必填列：upload_video_url")
        index = {name: position for position, name in enumerate(headers) if name}
        data_rows = [values for values in rows if any(value is not None and str(value).strip() for value in values)]
        if len(data_rows) > 3:
            raise HTTPException(400, "一次最多导入 3 条视频任务")

        for excel_row, values in enumerate(data_rows, start=2):
            def cell(column: str) -> Any:
                position = index.get(column)
                return values[position] if position is not None and position < len(values) else None
            try:
                source_url = validate_public_media_url(text_cell(cell("upload_video_url")), "upload_video_url")
                reference_urls = [
                    validate_public_media_url(item.strip(), "reference_image_url")
                    for item in re.split(r"[,，;；\n]+", text_cell(cell("reference_image_urls"))) if item.strip()
                ]
                if len(reference_urls) > 9:
                    raise ValueError("reference_image_urls 最多填写 9 个图片 URL")
                duration = probe_video_duration(source_url)
                aspect_ratio = text_cell(cell("aspect_ratio")) or "auto"
                quality = text_cell(cell("quality")) or "low"
                seed = int_cell(cell("seed"), None)
                task_id = insert_task(
                    prompt=text_cell(cell("prompt")), source_video=source_url,
                    reference_images=reference_urls, duration=duration, aspect_ratio=aspect_ratio,
                    seed=seed, quality=quality, display_name=media_display_name(source_url),
                )
                created.append(task_id)
            except Exception as exc:
                errors.append({"row": excel_row, "error": str(exc)})
        return {"created": len(created), "task_ids": created, "errors": errors}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Excel 导入失败：{exc}") from exc


@app.get("/api/tasks")
def list_tasks() -> dict[str, Any]:
    with connect_db() as conn:
        rows = conn.execute(
            "SELECT * FROM tasks WHERE deleted_at IS NULL ORDER BY created_at DESC, id DESC"
        ).fetchall()
        queued = [row["id"] for row in reversed(rows) if row["status"] == "queued"]
    positions = {task_id: i + 1 for i, task_id in enumerate(queued)}
    return {"items": [task_to_dict(row, positions.get(row["id"])) for row in rows]}


@app.delete("/api/tasks/{task_id}")
def delete_task(task_id: str) -> dict[str, Any]:
    with connect_db() as conn:
        row = conn.execute("SELECT * FROM tasks WHERE id=? AND deleted_at IS NULL", (task_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "任务不存在")
        now = utcnow()
        conn.execute(
            "UPDATE tasks SET cancel_requested=1,deleted_at=?,updated_at=? WHERE id=?",
            (now, now, task_id),
        )
    if row["status"] in {"starting", "running"}:
        try_cancel_comfy(row["engine_job_id"])
    if row["output_path"]:
        Path(row["output_path"]).unlink(missing_ok=True)
    return {"deleted": True, "was_running": row["status"] in {"starting", "running"}}


@app.post("/api/tasks/{task_id}/retry")
def retry_task(task_id: str) -> dict[str, Any]:
    with connect_db() as conn:
        row = conn.execute("SELECT status FROM tasks WHERE id=? AND deleted_at IS NULL", (task_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "任务不存在")
        if row["status"] not in {"failed", "cancelled"}:
            raise HTTPException(409, "只有失败或已取消的任务可以重试")
        conn.execute(
            "UPDATE tasks SET status='queued',progress=0,stage='等待队列',error=NULL,cancel_requested=0,"
            "engine_job_id=NULL,current_segment=0,output_width=NULL,output_height=NULL,"
            "started_at=NULL,finished_at=NULL,updated_at=? WHERE id=?",
            (utcnow(), task_id),
        )
    return {"id": task_id, "status": "queued"}


@app.get("/api/tasks/{task_id}/output")
def task_output(task_id: str):
    with connect_db() as conn:
        row = conn.execute(
            "SELECT name,status,output_path FROM tasks WHERE id=? AND deleted_at IS NULL", (task_id,)
        ).fetchone()
    if row is None:
        raise HTTPException(404, "任务不存在")
    if row["status"] != "completed" or not row["output_path"] or not Path(row["output_path"]).is_file():
        raise HTTPException(409, "成片尚未生成")
    filename = safe_name(row["name"]) + ".mp4"
    return FileResponse(row["output_path"], media_type="video/mp4", filename=filename)


@app.get("/api/template")
def excel_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "任务"
    headers = ["upload_video_url", "reference_image_urls", "prompt", "aspect_ratio", "quality", "seed"]
    sheet.append(headers)
    sheet.append([
        "https://your-domain.example/videos/source.mp4",
        "https://your-domain.example/images/person.jpg\nhttps://your-domain.example/images/product.png",
        "保持上传视频的场景和运镜，用参考图片替换人物和商品，包装文字尽量清晰。", "auto", "low", "",
    ])
    header_fill = PatternFill("solid", fgColor="253449")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor="FFF4CC")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[2].height = 38
    widths = [52, 62, 72, 18, 18, 18]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + i)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E2"
    sheet["A1"].comment = Comment(
        "必填；用户自行提供可公开访问的视频 URL。不能短于 4 秒；超过 15 秒会自动分段生成并合并。",
        "MiniMax H3 Studio",
    )
    sheet["B1"].comment = Comment("可选；用户自行提供人物或商品图片 URL，每行一个，最多 9 个。", "MiniMax H3 Studio")
    sheet["E1"].comment = Comment("可选；low、standard 或 high。默认 low，速度最快。", "MiniMax H3 Studio")
    ratio_validation = DataValidation(type="list", formula1='"auto,16:9,9:16,1:1,4:3,3:4,21:9"')
    sheet.add_data_validation(ratio_validation)
    ratio_validation.add("D2:D1000")
    quality_validation = DataValidation(type="list", formula1='"low,standard,high"')
    sheet.add_data_validation(quality_validation)
    quality_validation.add("E2:E1000")

    guide = workbook.create_sheet("字段说明")
    guide.append(["字段", "是否必填", "说明"])
    guide_rows = [
        ("upload_video_url", "是", "用户自行提供可公开访问的视频 URL。不能短于 4 秒；超过 15 秒会自动分段生成并合并。"),
        ("reference_image_urls", "否", "用户自行提供人物和商品参考图片 URL；每行一个，最多 9 个。"),
        ("prompt", "否", "说明每张参考图片用于替换人物还是商品，并填写其他保留要求。"),
        ("aspect_ratio", "否", "默认 auto，跟随上传视频比例；也可指定固定比例。"),
        ("quality", "否", "清晰度：low（低清，默认且最快）、standard（标清）或 high（高清且最慢）。"),
        ("seed", "否", "固定随机种子便于复现；空白时自动生成。"),
    ]
    for row in guide_rows:
        guide.append(row)
    for cell in guide[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 14
    guide.column_dimensions["C"].width = 72
    guide.freeze_panes = "A2"

    content = io.BytesIO()
    workbook.save(content)
    content.seek(0)
    headers_out = {"Content-Disposition": 'attachment; filename="h3_tasks_template.xlsx"'}
    return StreamingResponse(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_out,
    )


app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
